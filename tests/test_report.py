import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import report  # noqa: E402


def test_demo_report(tmp_path):
    csv_path = report.demo_csv(tmp_path / "s.csv")
    sales = report.read_csv(csv_path)
    out = report.build(sales, tmp_path / "r.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Сводка", "Менеджеры", "Товары", "Данные"]
    assert len(wb["Сводка"]._charts) == 2
    assert wb["Данные"].max_row == len([s for s in sales if s.day.year == max(x.day for x in sales).year]) + 1
    assert wb["Менеджеры"].max_row == 5


def test_bad_csv(tmp_path):
    p = tmp_path / "bad.csv"
    p.write_text("date;manager;category;product;qty;price\n2026-01-01;A;B;C;много;100\n", encoding="utf-8")
    try:
        report.read_csv(p)
    except SystemExit as e:
        assert "Строка 2" in str(e)
    else:
        raise AssertionError("ожидалась ошибка")
