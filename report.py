"""Автоматический отчёт о продажах: CSV-выгрузка из CRM → оформленный Excel с графиками.

Листы отчёта:
  «Сводка»    — ключевые цифры, выручка по месяцам (график) и доля категорий (диаграмма)
  «Менеджеры» — план/факт по менеджерам с условным форматированием
  «Товары»    — топ товаров по выручке
  «Данные»    — исходные строки с автофильтром

Запуск:
    python report.py --demo                      # сгенерировать пример данных и отчёт
    python report.py sales.csv --out report.xlsx
Формат CSV: date;manager;category;product;qty;price (разделитель «;», дата ГГГГ-ММ-ДД)
"""
from __future__ import annotations

import argparse
import csv
import random
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONTHS = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"]
HEAD_FILL = PatternFill("solid", fgColor="1F4E46")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9DEDB")
MONEY = '#,##0 "₽"'


@dataclass
class Sale:
    day: date
    manager: str
    category: str
    product: str
    qty: int
    price: float

    @property
    def total(self) -> float:
        return self.qty * self.price


def read_csv(path: Path) -> list[Sale]:
    sales = []
    with path.open(encoding="utf-8-sig", newline="") as f:
        for i, row in enumerate(csv.DictReader(f, delimiter=";"), start=2):
            try:
                sales.append(Sale(date.fromisoformat(row["date"]), row["manager"].strip(), row["category"].strip(),
                                  row["product"].strip(), int(row["qty"]), float(row["price"].replace(",", "."))))
            except (KeyError, ValueError) as exc:
                raise SystemExit(f"Строка {i}: не удалось разобрать ({exc}). Проверьте формат CSV.")
    if not sales:
        raise SystemExit("Файл пустой.")
    return sales


def demo_csv(path: Path, seed: int = 7) -> Path:
    rnd = random.Random(seed)
    catalog = {"Ноутбуки": [("Ноутбук 14", 89000), ("Ноутбук 16 Pro", 145000)], "Смартфоны": [("Смартфон X", 54000), ("Смартфон Lite", 21000)],
               "Аксессуары": [("Наушники", 12900), ("Чехол", 1500), ("Зарядка 65 Вт", 3900)], "Планшеты": [("Планшет 11", 39000)]}
    managers = ["Иванова А.", "Петров С.", "Сидоров К.", "Орлова Е."]
    start = date(date.today().year, 1, 1)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["date", "manager", "category", "product", "qty", "price"])
        for d in range((date.today() - start).days + 1):
            for _ in range(rnd.randint(2, 7)):
                cat = rnd.choice(list(catalog))
                prod, price = rnd.choice(catalog[cat])
                w.writerow([(start + timedelta(days=d)).isoformat(), rnd.choice(managers), cat, prod, rnd.randint(1, 3), price])
    return path


def style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, values: list[int]) -> None:
    for i, w in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(sales: list[Sale], out: Path, plan_per_month: float = 2_500_000) -> Path:
    wb = Workbook()
    year = max(s.day for s in sales).year
    sales = [s for s in sales if s.day.year == year]

    # ---- Сводка ----
    ws = wb.active
    ws.title = "Сводка"
    total = sum(s.total for s in sales)
    orders = len(sales)
    ws["A1"] = f"Отчёт о продажах за {year} год"
    ws["A1"].font = Font(size=18, bold=True, color="1F4E46")
    ws["A2"] = f"Данные с {min(s.day for s in sales):%d.%m.%Y} по {max(s.day for s in sales):%d.%m.%Y}"
    ws["A2"].font = Font(italic=True, color="6B7280")
    kpis = [("Выручка", total, MONEY), ("Продаж", orders, "#,##0"), ("Средний чек", total / orders, MONEY),
            ("Товаров продано", sum(s.qty for s in sales), "#,##0")]
    for i, (label, value, fmt) in enumerate(kpis):
        col = 1 + i * 2
        ws.cell(row=4, column=col, value=label).font = Font(color="6B7280")
        cell = ws.cell(row=5, column=col, value=value)
        cell.font, cell.number_format = Font(size=16, bold=True), fmt

    by_month = defaultdict(float)
    for s in sales:
        by_month[s.day.month] += s.total
    ws["A8"], ws["B8"], ws["C8"], ws["D8"] = "Месяц", "Выручка", "План", "Выполнение"
    style_header(ws, 8, 4)
    last_month = max(by_month)
    for m in range(1, last_month + 1):
        r = 8 + m
        ws.cell(row=r, column=1, value=MONTHS[m - 1])
        ws.cell(row=r, column=2, value=round(by_month.get(m, 0))).number_format = MONEY
        ws.cell(row=r, column=3, value=plan_per_month).number_format = MONEY
        pct = ws.cell(row=r, column=4, value=f"=B{r}/C{r}")
        pct.number_format = "0%"
    end = 8 + last_month
    ws.conditional_formatting.add(f"D9:D{end}", CellIsRule(operator="lessThan", formula=["0.9"], font=Font(color="B42318", bold=True)))
    ws.conditional_formatting.add(f"D9:D{end}", CellIsRule(operator="greaterThanOrEqual", formula=["1"], font=Font(color="15803D", bold=True)))

    chart = BarChart()
    chart.title, chart.y_axis.title = "Выручка по месяцам", "₽"
    chart.add_data(Reference(ws, min_col=2, min_row=8, max_row=end), titles_from_data=True)
    chart.add_data(Reference(ws, min_col=3, min_row=8, max_row=end), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=9, max_row=end))
    chart.height, chart.width = 8, 18
    ws.add_chart(chart, "F8")

    by_cat = defaultdict(float)
    for s in sales:
        by_cat[s.category] += s.total
    start_row = end + 3
    ws.cell(row=start_row, column=1, value="Категория")
    ws.cell(row=start_row, column=2, value="Выручка")
    style_header(ws, start_row, 2)
    for i, (cat, v) in enumerate(sorted(by_cat.items(), key=lambda kv: -kv[1]), 1):
        ws.cell(row=start_row + i, column=1, value=cat)
        ws.cell(row=start_row + i, column=2, value=round(v)).number_format = MONEY
    pie = PieChart()
    pie.title = "Доля категорий"
    pie.add_data(Reference(ws, min_col=2, min_row=start_row, max_row=start_row + len(by_cat)), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=start_row + 1, max_row=start_row + len(by_cat)))
    pie.height, pie.width = 8, 12
    ws.add_chart(pie, f"F{start_row}")
    widths(ws, [18, 16, 16, 14])

    # ---- Менеджеры ----
    wm = wb.create_sheet("Менеджеры")
    headers = ["Менеджер", "Продаж", "Выручка", "Средний чек", "План", "Выполнение"]
    wm.append(headers)
    style_header(wm, 1, len(headers))
    per_manager = defaultdict(list)
    for s in sales:
        per_manager[s.manager].append(s)
    manager_plan = plan_per_month * last_month / max(1, len(per_manager))
    for r, (name, items) in enumerate(sorted(per_manager.items(), key=lambda kv: -sum(x.total for x in kv[1])), start=2):
        rev = sum(x.total for x in items)
        wm.append([name, len(items), round(rev), f"=C{r}/B{r}", round(manager_plan), f"=C{r}/E{r}"])
        for c, fmt in ((3, MONEY), (4, MONEY), (5, MONEY), (6, "0%")):
            wm.cell(row=r, column=c).number_format = fmt
    last = 1 + len(per_manager)
    wm.conditional_formatting.add(f"C2:C{last}", DataBarRule(start_type="min", end_type="max", color="7BC4B2"))
    wm.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="lessThan", formula=["1"], fill=PatternFill("solid", fgColor="FDE2E1")))
    wm.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=PatternFill("solid", fgColor="DCFCE7")))
    widths(wm, [20, 10, 16, 14, 16, 13])

    # ---- Товары ----
    wt = wb.create_sheet("Товары")
    wt.append(["Товар", "Категория", "Штук", "Выручка", "Доля"])
    style_header(wt, 1, 5)
    per_product = defaultdict(lambda: [0, 0.0, ""])
    for s in sales:
        p = per_product[s.product]
        p[0] += s.qty; p[1] += s.total; p[2] = s.category
    for r, (name, (qty, rev, cat)) in enumerate(sorted(per_product.items(), key=lambda kv: -kv[1][1]), start=2):
        wt.append([name, cat, qty, round(rev), rev / total])
        wt.cell(row=r, column=4).number_format = MONEY
        wt.cell(row=r, column=5).number_format = "0.0%"
    widths(wt, [22, 16, 10, 16, 10])

    # ---- Данные ----
    wd = wb.create_sheet("Данные")
    wd.append(["Дата", "Менеджер", "Категория", "Товар", "Кол-во", "Цена", "Сумма"])
    style_header(wd, 1, 7)
    for r, s in enumerate(sorted(sales, key=lambda x: x.day), start=2):
        wd.append([s.day, s.manager, s.category, s.product, s.qty, s.price, f"=E{r}*F{r}"])
        wd.cell(row=r, column=1).number_format = "DD.MM.YYYY"
        wd.cell(row=r, column=6).number_format = MONEY
        wd.cell(row=r, column=7).number_format = MONEY
    wd.auto_filter.ref = wd.dimensions
    wd.freeze_panes = "A2"
    widths(wd, [12, 18, 14, 18, 8, 12, 14])

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row):
            for cell in row:
                if cell.value is not None and sheet.title != "Сводка":
                    cell.border = Border(bottom=THIN)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description="Excel-отчёт о продажах из CSV")
    ap.add_argument("csv", nargs="?", type=Path, help="выгрузка продаж")
    ap.add_argument("--demo", action="store_true", help="создать пример данных")
    ap.add_argument("--plan", type=float, default=2_500_000, help="план выручки в месяц")
    ap.add_argument("--out", type=Path, default=Path("output/report.xlsx"))
    args = ap.parse_args()
    if args.demo:
        args.csv = demo_csv(Path("output/demo_sales.csv"))
        print(f"Пример данных: {args.csv}")
    if not args.csv:
        ap.error("укажите CSV-файл или --demo")
    out = build(read_csv(args.csv), args.out, args.plan)
    print(f"Отчёт готов: {out}")


if __name__ == "__main__":
    main()
